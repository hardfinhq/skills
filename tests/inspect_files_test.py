# Copyright (c) 2026. Hardfin, Inc. All rights reserved.
"""Test dependency-free worksheet value exports."""

import csv
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font
from openpyxl.utils.datetime import CALENDAR_MAC_1904


ROOT = Path(__file__).resolve().parent
SCRIPT = ROOT.parent / "skills/hardfin-import/scripts/inspect_files.py"
SCRIPT_IN_GIT = "skills/hardfin-import/scripts/inspect_files.py"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
PACKAGE_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SHEET_NAME = "Site / west"
CSV_NAME = "Fixture - Site _ west.csv"
LONG_IDENTIFIER = "123456789012345678901"
SERIAL = 45292
OUT_OF_RANGE_SERIAL = 4050123456789
LAST_COLUMN = 27


def tag(name):
    """Return a spreadsheet XML name."""
    return "{%s}%s" % (MAIN_NS, name)


def make_workbook(path, date1904=False):
    """Create a workbook with representative stored cell types."""
    workbook = Workbook()
    if date1904:
        workbook.epoch = CALENDAR_MAC_1904

    sheet = workbook.active
    sheet.title = "Site west"
    sheet.append(["placeholder", "placeholder", 1, SERIAL, SERIAL + 0.5, True, "#REF!", "=1+1", False,
                  4050123456789, 1234567890123456, 1])
    sheet["D1"].number_format = "yyyy-mm-dd"
    sheet["M1"] = OUT_OF_RANGE_SERIAL
    sheet["M1"].number_format = "yyyy-mm-dd"
    sheet["E1"].number_format = "yyyy-mm-dd hh:mm:ss"
    sheet["A3"] = "Hidden value"
    sheet.row_dimensions[3].hidden = True
    sheet["B5"] = "Tail"
    sheet["AA5"] = "Last column"
    for row, serial in enumerate((59, 60, 61, 0, 0.5), 8):
        sheet.cell(row, 1, serial).number_format = "yyyy-mm-dd hh:mm:ss"

    workbook.create_sheet("Empty")
    workbook.save(path)

    # XML fixtures include Excel data that openpyxl cannot save, such as cached formulas and slash-containing names.
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    cells = {node.attrib["r"]: node for node in root.iter(tag("c"))}
    for reference, value in (("C1", LONG_IDENTIFIER), ("E1", "45292.52425925926"),
                             ("H1", "2"), ("L1", "1.234567890123456E+20")):
        cells[reference].find(tag("v")).text = value

    inline = cells["A1"].find(tag("is"))
    inline.clear()
    for text in ("Café, ", '"inline"\ntext'):
        ET.SubElement(ET.SubElement(inline, tag("r")), tag("t")).text = text

    cells["B1"].clear()
    cells["B1"].attrib.update(r="B1", t="s")
    ET.SubElement(cells["B1"], tag("v")).text = "0"
    shared = ET.Element(tag("sst"))
    item = ET.SubElement(shared, tag("si"))
    for text in ("Shared ", "rich text"):
        ET.SubElement(ET.SubElement(item, tag("r")), tag("t")).text = text

    parts["xl/sharedStrings.xml"] = ET.tostring(shared)
    relations = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    ET.SubElement(relations, "{%s}Relationship" % PACKAGE_NS,
                  Id="shared", Target="sharedStrings.xml", Type=REL_NS + "/sharedStrings")
    parts["xl/_rels/workbook.xml.rels"] = ET.tostring(relations)
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    root = ET.fromstring(parts["xl/workbook.xml"])
    root.find(tag("sheets"))[0].set("name", SHEET_NAME)
    parts["xl/workbook.xml"] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)


def run_inspector(*arguments, source=None):
    """Run the inspector in an isolated Python process."""
    command = [sys.executable, "-I", str(SCRIPT) if source is None else "-"]
    return subprocess.run(command + list(map(str, arguments)), input=source, text=True,
                          capture_output=True, check=True).stdout


class InspectFilesTest(unittest.TestCase):
    """Check exported values and unchanged reports."""

    @classmethod
    def setUpClass(cls):
        """Prepare workbooks and their exported records."""
        cls.temporary = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temporary.cleanup)
        cls.folder = Path(cls.temporary.name)
        cls.workbook = cls.folder / "Fixture.xlsx"
        make_workbook(cls.workbook)
        cls.output = cls.folder / "values"
        cls.report = json.loads(run_inspector("--json", "--values", cls.output, cls.workbook))
        with (cls.output / CSV_NAME).open(encoding="utf-8", newline="") as stream:
            cls.records = list(csv.reader(stream))

        cls.rows = {int(row[0]): row[1:] for row in cls.records[1:]}

    def test_inline_strings(self):
        self.assertEqual(self.rows[1][0], 'Café, "inline"\ntext')

    def test_shared_strings(self):
        self.assertEqual(self.rows[1][1], "Shared rich text")

    def test_exact_numbers(self):
        self.assertEqual(self.rows[1][2], LONG_IDENTIFIER)
        self.assertEqual(self.rows[1][9:12], ["4050123456789", "1234567890123456", "1.234567890123456E+20"])

    def test_1900_date(self):
        self.assertEqual(self.rows[1][3], "2024-01-01")

    def test_out_of_range_date(self):
        self.assertEqual(self.rows[1][12], str(OUT_OF_RANGE_SERIAL))

    def test_1904_date(self):
        path = self.folder / "Mac.xlsm"
        make_workbook(path, date1904=True)
        run_inspector("--values", self.output, path)
        with (self.output / "Mac - Site _ west.csv").open(encoding="utf-8", newline="") as stream:
            rows = list(csv.reader(stream))

        self.assertEqual(rows[1][4], "2028-01-02")

    def test_time_part(self):
        self.assertEqual(self.rows[1][4], "2024-01-01T12:34:56")
        self.assertEqual(self.rows[12][0], "1899-12-31T12:00:00")

    def test_leap_bug(self):
        self.assertEqual([self.rows[row][0] for row in range(8, 12)],
                         ["1900-02-28", "1900-02-29", "1900-03-01", "1899-12-31"])

    def test_boolean(self):
        self.assertEqual(self.rows[1][5], "TRUE")
        self.assertEqual(self.rows[1][8], "FALSE")

    def test_error(self):
        self.assertEqual(self.rows[1][6], "#REF!")

    def test_cached_formula(self):
        self.assertEqual(self.rows[1][7], "2")

    def test_hidden_row(self):
        self.assertEqual(self.rows[3][0], "Hidden value")

    def test_blank_row(self):
        self.assertNotIn(2, self.rows)
        self.assertNotIn(4, self.rows)

    def test_row_numbers(self):
        self.assertEqual([row[0] for row in self.records], ["row", "1", "3", "5", "8", "9", "10", "11", "12"])

    def test_columns(self):
        self.assertEqual(self.records[0], ["row"] + list("ABCDEFGHIJKLMNOPQRSTUVWXYZ") + ["AA"])
        self.assertTrue(all(len(row) == LAST_COLUMN + 1 for row in self.records))
        self.assertEqual(self.rows[5][0:2], ["", "Tail"])
        self.assertEqual(self.rows[5][-1], "Last column")

    def test_safe_name_and_json_path(self):
        sheets = self.report["files"][0]["sheets"]
        self.assertEqual(sheets[0]["csv_path"], str(self.output / CSV_NAME))
        self.assertEqual(sheets[0]["name"], SHEET_NAME)
        self.assertTrue(Path(sheets[0]["csv_path"]).is_file())

    def test_empty_sheet(self):
        with (self.output / "Fixture - Empty.csv").open(encoding="utf-8", newline="") as stream:
            self.assertEqual(list(csv.reader(stream)), [["row"]])

    def test_text_report(self):
        output = run_inspector("--values", self.output, self.workbook)
        lines = [line for line in output.splitlines() if "CSV files written:" in line]
        self.assertEqual(len(lines), 1)
        self.assertIn(str(self.output / CSV_NAME), lines[0])
        self.assertIn(str(self.output / "Fixture - Empty.csv"), lines[0])

    def test_csv_unchanged(self):
        path = self.folder / "source.csv"
        content = b"name,value\r\nTest,123\r\n"
        path.write_bytes(content)
        output = self.folder / "csv-only"
        report = json.loads(run_inspector("--json", "--values", output, path))
        self.assertEqual(list(output.iterdir()), [])
        self.assertEqual(path.read_bytes(), content)
        self.assertNotIn("csv_path", report["files"][0]["sheets"][0])

    def test_headerless_grid(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Grid"
        sheet["E1"] = "Equipment allocation"
        sheet.append(["AB-1", "Analyser", "2026-02-19", "Main store", None, "FUNCTIONAL", 100, 10, 36])
        path = self.folder / "grid.xlsx"
        workbook.save(path)
        report = json.loads(run_inspector("--json", path))["files"][0]
        self.assertIsNone(report["sheets"][0]["header_row"])

    def test_defined_names(self):
        from openpyxl.workbook.defined_name import DefinedName
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Register"
        sheet.append(["Serial", "Name"])
        sheet.append(["AB-1", "First"])
        sheet.append(["AB-2", "Second"])
        workbook.defined_names.add(DefinedName("ActiveUnits", attr_text="Register!$A$2:$B$2"))
        workbook.defined_names.add(DefinedName("_SpareUnits", attr_text="Register!$A$3:$B$3"))
        workbook.defined_names.add(DefinedName("_xlnm._FilterDatabase", localSheetId=0, hidden=True,
                                               attr_text="Register!$A$1:$B$3"))
        path = self.folder / "named.xlsx"
        workbook.save(path)
        findings = json.loads(run_inspector("--json", path))["files"][0]["findings"]
        named = [item for item in findings if item["kind"] == "defined_names"]
        self.assertEqual(len(named), 1)
        self.assertEqual(named[0]["entries"], ['"ActiveUnits": Register!$A$2:$B$2', '"_SpareUnits": Register!$A$3:$B$3'])

    def test_truncated_csv(self):
        path = self.folder / "cut.csv"
        path.write_bytes(b"serial,name\nAB-1,First\nAB-2,\"Second, cut")
        report = json.loads(run_inspector("--json", path))["files"][0]
        self.assertEqual(report["status"], "inspected")
        findings = {item["kind"]: item["detail"] for item in report["sheets"][0]["findings"]}
        self.assertIn("Reading stopped at line 3", findings["truncated"])
        self.assertEqual(report["sheets"][0]["used_range"], "A1:B2")

    def test_utf16_delimited(self):
        for codec, bom, encoding in (("utf-16-le", b"\xff\xfe", "UTF-16LE with BOM"),
                                     ("utf-16-be", b"\xfe\xff", "UTF-16BE with BOM")):
            with self.subTest(encoding=encoding), tempfile.TemporaryDirectory() as temporary:
                folder = Path(temporary)
                path = folder / "Equipment.csv"
                path.write_bytes(bom + "Serial\tName\r\nAB-1\tCafé\r\n".encode(codec))
                report = json.loads(run_inspector("--json", "--values", folder / "values", path))
                sheet = report["files"][0]["sheets"][0]
                findings = {item["kind"]: item["detail"] for item in sheet["findings"]}
                self.assertEqual(findings["encoding"], encoding)
                self.assertEqual(findings["delimiter"], 'Delimiter "\\t" (csv.Sniffer)')
                with Path(sheet["csv_path"]).open(encoding="utf-8", newline="") as stream:
                    self.assertEqual(list(csv.reader(stream)),
                                     [["row", "A", "B"], ["1", "Serial", "Name"], ["2", "AB-1", "Café"]])

    def test_zip_corrupt_member_reported_alone(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "Exports.zip"
            with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as archive:
                archive.writestr("broken.csv", "Serial,Name\nCORRUPT-1,First\n")
                archive.writestr("assets.csv", "Serial,Name\nAB-1,First\n")

            data = path.read_bytes()
            path.write_bytes(data.replace(b"CORRUPT-1", b"CORRUPT-2", 1))
            report = json.loads(run_inspector("--json", path))["files"][0]
            statuses = {Path(member["path"]).name: member["status"] for member in report["members"]}
            self.assertEqual(statuses, {"broken.csv": "error", "assets.csv": "inspected"})

    def test_zip_members(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            path = folder / "Exports.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("nested/assets.csv", "Serial,Name\nAB-1,First\n")
                archive.writestr("nested/Equipment.xlsx", self.workbook.read_bytes())
                archive.writestr("__MACOSX/._x.csv", b"metadata")
                archive.writestr(".DS_Store", b"metadata")
                archive.writestr("nested/._other.xlsx", b"metadata")
                archive.writestr("../escape.txt", "Name\nSafe\n")
                archive.writestr("unsupported.exe", b"never execute")

            before = set(folder.rglob("*"))
            report = json.loads(run_inspector("--json", path))["files"][0]
            self.assertEqual(set(folder.rglob("*")), before)
            self.assertEqual(report["status"], "inspected")
            self.assertEqual([member["path"] for member in report["members"]],
                             [str(path) + "/" + name for name in
                              ("nested/assets.csv", "nested/Equipment.xlsx", "../escape.txt")])
            self.assertTrue(all(member["status"] == "inspected" for member in report["members"]))
            metadata = report["findings"]
            self.assertEqual(len(metadata), 1)
            self.assertIn("metadata holds no records", metadata[0]["detail"])
            self.assertEqual(metadata[0]["entries"],
                             ['"__MACOSX/._x.csv"', '".DS_Store"', '"nested/._other.xlsx"'])
            self.assertIn("Hidden rows", str(report["members"][1]["sheets"][0]["findings"]))
            output = folder / "values"
            report = json.loads(run_inspector("--json", "--values", output, path))["files"][0]
            exports = [Path(sheet["csv_path"]) for member in report["members"] for sheet in member["sheets"]]
            self.assertEqual(len(exports), 4)
            self.assertEqual(set(folder.rglob("*")) - before, {output, *exports})
            self.assertTrue(all(file.parent == output and file.name.startswith("Exports.zip_") for file in exports))
            self.assertEqual(exports[0].name, "Exports.zip_nested_assets - assets.csv.csv")
            with exports[0].open(encoding="utf-8", newline="") as stream:
                self.assertEqual(list(csv.reader(stream)),
                                 [["row", "A", "B"], ["1", "Serial", "Name"], ["2", "AB-1", "First"]])

    def test_invisible_characters(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            values = ["Right\u200f", "\tLeading", "Line\nfeed", "Line\r\nbreak", "\rEdge", "Edge\r"]
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Value"])
            for value in values:
                sheet.append([value])

            xlsx = folder / "Controls.xlsx"
            workbook.save(xlsx)
            # Literal carriage returns need XML character references to survive XML newline normalization.
            with zipfile.ZipFile(xlsx) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}

            with zipfile.ZipFile(xlsx, "w") as archive:
                for name, content in parts.items():
                    archive.writestr(name, content.replace(b"\r", b"&#13;"))

            text = folder / "Controls.csv"
            with text.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerows([["Value"]] + [[value] for value in values])

            utf16 = folder / "Controls-utf16.csv"
            utf16.write_bytes(text.read_bytes().decode("utf-8").encode("utf-16"))
            for path in (text, xlsx, utf16):
                with self.subTest(extension=path.suffix):
                    report = json.loads(run_inspector("--json", "--values", folder / "values", path))
                    sheet = report["files"][0]["sheets"][0]
                    findings = [item for item in sheet["findings"] if item["kind"] == "invisible_characters"]
                    self.assertEqual(len(findings), 1)
                    self.assertEqual(findings[0]["entries"], ['A2 ("Value"): U+200F', 'A3 ("Value"): U+0009',
                                                            'A6 ("Value"): U+000D', 'A7 ("Value"): U+000D'])
                    if "csv_path" in sheet:
                        with Path(sheet["csv_path"]).open(encoding="utf-8", newline="") as stream:
                            self.assertEqual([row[1] for row in list(csv.reader(stream))[2:]], values)

    def test_partial_strikethrough(self):
        for storage in ("inline", "shared"):
            with self.subTest(storage=storage), tempfile.TemporaryDirectory() as temporary:
                folder = Path(temporary)
                path = folder / "Rich.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["Item"])
                sheet["A2"] = CellRichText(TextBlock(InlineFont(strike=True), "Old name"),
                                           TextBlock(InlineFont(strike=False), "New name"), " suffix")
                sheet["A3"] = "Whole cell"
                sheet["A3"].font = Font(strike=True)
                sheet["A4"] = CellRichText(TextBlock(InlineFont(strike=True), "All struck"))
                sheet["A5"] = CellRichText(TextBlock(InlineFont(strike=False), "Not struck"), " plain")
                workbook.save(path)
                with zipfile.ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}

                root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
                cell = next(node for node in root.iter(tag("c")) if node.get("r") == "A2")
                inline = cell.find(tag("is"))
                inline.find(tag("r") + "/" + tag("rPr") + "/" + tag("strike")).attrib.clear()
                # Shared strings use the same rich runs in a separate part referenced by the cell.
                if storage == "shared":
                    cell.remove(inline)
                    cell.set("t", "s")
                    ET.SubElement(cell, tag("v")).text = "0"
                    inline.tag = tag("si")
                    shared = ET.Element(tag("sst"))
                    shared.append(inline)
                    parts["xl/sharedStrings.xml"] = ET.tostring(shared)
                    relations = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
                    ET.SubElement(relations, "{%s}Relationship" % PACKAGE_NS, Id="shared",
                                  Target="sharedStrings.xml", Type=REL_NS + "/sharedStrings")
                    parts["xl/_rels/workbook.xml.rels"] = ET.tostring(relations)

                parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
                with zipfile.ZipFile(path, "w") as archive:
                    for name, content in parts.items():
                        archive.writestr(name, content)

                report = json.loads(run_inspector("--json", "--values", folder / "values", path))
                sheet = report["files"][0]["sheets"][0]
                findings = [item for item in sheet["findings"] if item["kind"] == "partial_strikethrough"]
                self.assertEqual(len(findings), 1)
                self.assertEqual(findings[0]["entries"],
                                 ['A2 ("Item"): struck "Old name"; not struck "New name", " suffix"'])
                whole = [item for item in sheet["findings"] if item["kind"] == "strikethrough"]
                self.assertEqual(len(whole), 1)
                self.assertEqual(whole[0]["entries"], ['A3 ("Item"): value "Whole cell"'])
                with Path(sheet["csv_path"]).open(encoding="utf-8", newline="") as stream:
                    self.assertEqual(list(csv.reader(stream))[2], ["2", "Old nameNew name suffix"])

    def test_without_values_unchanged(self):
        source = subprocess.run(["git", "show", "HEAD:" + SCRIPT_IN_GIT], cwd=ROOT,
                                capture_output=True, text=True, check=True).stdout
        for flags in ((), ("--json",)):
            self.assertEqual(run_inspector(*flags, self.workbook),
                             run_inspector(*flags, self.workbook, source=source))


if __name__ == "__main__":
    unittest.main()
